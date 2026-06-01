from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width  = Inches(13.33)
prs.slide_height = Inches(7.5)

def rgb(r, g, b): return RGBColor(r, g, b)

def add_rect(slide, l, t, w, h, fill_rgb=None, border_rgb=None, border_pt=0):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    if fill_rgb:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_rgb
    else:
        shape.fill.background()
    if border_rgb and border_pt:
        shape.line.color.rgb = border_rgb
        shape.line.width = Pt(border_pt)
    else:
        shape.line.fill.background()
    return shape

def add_text(slide, l, t, w, h, text, size, bold=False,
             color=rgb(0,0,0), align=PP_ALIGN.LEFT):
    txBox = slide.shapes.add_textbox(l, t, w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return txBox

W = prs.slide_width
H = prs.slide_height
pad = Inches(0.5)
hdr_h = Inches(1.55)
card_w = W - 2 * pad

all_topics = [
    ("01", "Werbewirkungsforschung", rgb(160,100,0), rgb(255,243,205),
     "Werbewirkungsmessung einer NEOH TV-Kampagne auf dem deutschen Markt",
     "Empirische Analyse der Werbewirkung einer NEOH Fernsehwerbekampagne in Deutschland — "
     "inkl. Messung von Bekanntheit, Einstellungsänderung und Kaufabsicht anhand geeigneter "
     "Wirkungsmodelle und Methoden der Marktforschung."),
    ("02", "Marketing Analytics", rgb(160,100,0), rgb(255,243,205),
     "Marketing Mix Modelling für NEOH",
     "Entwicklung und Anwendung eines Marketing Mix Models zur quantitativen Analyse der Treiber "
     "von Absatz und Umsatz bei NEOH — inkl. Dekomposition des Mediamix, ROI-Messung einzelner "
     "Kanäle, Handlungsempfehlungen für die Budgetallokation sowie Länderbudgetallokation."),
    ("03", "AI & Marketing", rgb(26,111,168), rgb(232,244,253),
     "Artificial Intelligence in (International) Marketing & Business",
     "Empirische Untersuchung des Einsatzes und der Wirkung von KI-Technologien im Marketing "
     "und internationalen Geschäftsumfeld — z. B. Personalisierung, Automatisierung, generative "
     "KI oder KI-gestützte Entscheidungsprozesse."),
    ("04", "Brand Growth", rgb(30,126,70), rgb(234,247,239),
     "How Small Brands Grow (Internationally)",
     "Empirische Analyse der Wachstumsstrategien kleiner Marken im nationalen und internationalen "
     "Kontext — unter Berücksichtigung von Markenbekanntheit, Distributionsaufbau, Positionierung "
     "und den Gesetzmäßigkeiten des empirischen Marketings (z. B. Ehrenberg-Bass)."),
    ("05", "Brand Communications", rgb(180,60,60), rgb(255,235,235),
     "(International) Brand Communications: Brand Engagement & Brand Communities",
     "Empirische Untersuchung internationaler Markenkommunikation mit Fokus auf Brand Engagement "
     "und Brand Communities — z. B. Aufbau und Pflege von Markengemeinschaften, Engagement-Treiber "
     "und deren Wirkung auf Markenloyalität und -identifikation."),
    ("06", "Replikation", rgb(107,63,160), rgb(243,238,255),
     "Replikationsstudien",
     "Empirische Replikation bestehender Studien aus dem Bereich Marketing oder internationales "
     "Business — zur Überprüfung der Generalisierbarkeit von Befunden in neuen Kontexten, "
     "Märkten oder Zeiträumen."),
]

def build_slide(prs, topics, slide_num, total_slides, show_hints=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Header
    add_rect(slide, 0, 0, W, hdr_h, fill_rgb=rgb(26, 26, 46))
    add_text(slide, pad, Inches(0.16), W - pad*2, Inches(0.52),
             "Masterarbeitsthemen", 36, bold=True, color=rgb(255,255,255))
    add_text(slide, pad, Inches(0.66), W - pad*2, Inches(0.3),
             "Betreute Abschlussarbeiten · Ausschreibung 2026", 15,
             color=rgb(180,180,200))

    # NEOH badge (only on slide 1)
    if slide_num == 1:
        badge_t = Inches(1.04)
        add_rect(slide, pad, badge_t, Inches(4.4), Inches(0.36),
                 fill_rgb=rgb(50,50,80), border_rgb=rgb(100,100,140), border_pt=0.5)
        add_text(slide, pad + Inches(0.15), badge_t + Inches(0.06), Inches(4.2), Inches(0.28),
                 "✦  Themen 1 & 2 in Kooperation mit  NEOH", 12,
                 color=rgb(232,197,71))

    # Page indicator
    add_text(slide, W - Inches(1.2), Inches(0.1), Inches(1.0), Inches(0.3),
             f"{slide_num} / {total_slides}", 10,
             color=rgb(140,140,170), align=PP_ALIGN.RIGHT)

    # Cards
    n = len(topics)
    footer_h = Inches(0.3)
    avail_h = H - hdr_h - footer_h - Inches(0.15)
    if show_hints:
        avail_h -= Inches(0.62)
    card_gap = Inches(0.1)
    card_h = (avail_h - card_gap * (n - 1)) / n
    card_top = hdr_h + Inches(0.12)

    for i, (num, tag_text, tag_fg, tag_bg, title, desc) in enumerate(topics):
        t = card_top + i * (card_h + card_gap)

        add_rect(slide, pad, t, card_w, card_h,
                 fill_rgb=rgb(255,255,255), border_rgb=rgb(220,225,235), border_pt=0.75)

        # Number
        add_text(slide, pad + Inches(0.1), t + Inches(0.1), Inches(0.5), card_h,
                 num, 32, bold=True, color=rgb(215,215,215))

        # Tag
        tag_t = t + Inches(0.1)
        tag_w = Inches(2.1)
        tag_pill_h = Inches(0.26)
        add_rect(slide, pad + Inches(0.62), tag_t, tag_w, tag_pill_h, fill_rgb=tag_bg)
        add_text(slide, pad + Inches(0.72), tag_t + Inches(0.04), tag_w, tag_pill_h,
                 tag_text.upper(), 10, bold=True, color=tag_fg)

        # Title
        add_text(slide, pad + Inches(0.62), t + Inches(0.38), card_w - Inches(0.78),
                 Inches(0.36), title, 16, bold=True, color=rgb(26,26,46))

        # Description
        add_text(slide, pad + Inches(0.62), t + Inches(0.72), card_w - Inches(0.78),
                 card_h - Inches(0.78), desc, 12, color=rgb(90,106,128))

    # Hints (last slide only)
    if show_hints:
        hint_top = H - footer_h - Inches(0.6)
        add_rect(slide, pad, hint_top, card_w, Inches(0.54),
                 fill_rgb=rgb(247,248,250), border_rgb=rgb(228,232,238), border_pt=0.75)
        add_text(slide, pad + Inches(0.2), hint_top + Inches(0.05), card_w, Inches(0.18),
                 "ALLGEMEINE HINWEISE", 10, bold=True, color=rgb(138,150,168))
        add_text(slide, pad + Inches(0.2), hint_top + Inches(0.24), card_w, Inches(0.28),
                 "✦  Es werden nur empirische Arbeiten betreut (qualitativ oder quantitativ).    "
                 "✦  Die Arbeiten können auf Deutsch oder Englisch verfasst werden.", 12,
                 color=rgb(74,85,104))

    # Footer
    footer_t = H - footer_h
    add_rect(slide, 0, footer_t, W, footer_h, fill_rgb=rgb(247,248,250))
    add_text(slide, pad, footer_t + Inches(0.06), Inches(5), footer_h,
             "Kooperationspartner: NEOH GmbH", 9, color=rgb(170,170,170))
    add_text(slide, W - Inches(2), footer_t + Inches(0.06), Inches(1.8), footer_h,
             "Stand: Mai 2026", 9, color=rgb(170,170,170), align=PP_ALIGN.RIGHT)

# Slide 1: topics 1–3
build_slide(prs, all_topics[:3], slide_num=1, total_slides=2, show_hints=False)
# Slide 2: topics 4–6 + hints
build_slide(prs, all_topics[3:], slide_num=2, total_slides=2, show_hints=True)

prs.save("/home/user/My-first-repo/masterarbeitsthemen.pptx")
print("Done")
